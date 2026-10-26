"""Any-file attachment handling with text extraction and vision payloads.

Extensions are no longer restricted.  Common office formats, source code,
PDF (including scanned pages), images, archives and arbitrary binaries are
accepted.  Text is extracted when possible; images and scanned pages are also
exposed to vision-capable models.
"""

from __future__ import annotations

import base64
import binascii
import io
import json
import re
import uuid
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from . import config

MAX_BYTES = 60 * 1024 * 1024
MAX_CHARS = 120000
MAX_IMAGES = 4
MAX_IMAGE_SIDE = 1600

TEXT_SUFFIXES = {
    ".txt", ".md", ".markdown", ".csv", ".tsv", ".json", ".jsonl", ".xml", ".yaml",
    ".yml", ".toml", ".ini", ".cfg", ".conf", ".log", ".rst", ".tex", ".bib",
    ".py", ".pyw", ".m", ".mlx", ".c", ".h", ".cc", ".cpp", ".cxx", ".hpp",
    ".java", ".kt", ".scala", ".go", ".rs", ".js", ".jsx", ".ts", ".tsx", ".vue",
    ".css", ".scss", ".less", ".html", ".htm", ".php", ".rb", ".pl", ".lua",
    ".sh", ".bash", ".ps1", ".bat", ".cmd", ".sql", ".r", ".jl", ".v", ".sv",
    ".vhd", ".vhdl", ".asm", ".s", ".m4", ".cmake", ".mk", ".gradle", ".proto",
    ".srt", ".ass", ".vtt", ".diff", ".patch", ".gitignore", ".env",
}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}
AUDIO_SUFFIXES = {".mp3", ".wav", ".m4a", ".flac", ".ogg", ".aac"}
VIDEO_SUFFIXES = {".mp4", ".mov", ".webm", ".mkv", ".avi"}
ARCHIVE_SUFFIXES = {".zip", ".7z", ".rar", ".tar", ".gz", ".bz2", ".xz"}
LEGACY_OFFICE = {".doc", ".ppt", ".xls"}
_INFO_CACHE: Dict[str, Dict[str, Any]] = {}
_CACHE_ORDER: List[str] = []


def _decode_text(raw: bytes) -> str:
    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _ole_strings(raw: bytes, limit: int = 30000) -> str:
    """Best-effort text recovery from legacy OLE .doc/.ppt/.xls files."""
    chunks: List[str] = []
    # UTF-16LE runs (Word/PowerPoint store most text this way).
    try:
        decoded = raw.decode("utf-16-le", errors="ignore")
        for run in re.findall(r"[\u4e00-\u9fffA-Za-z0-9 \t,.;:+\-*/=()\[\]{}<>_%°µΩ]{4,}", decoded):
            run = run.strip()
            if len(run) >= 4 and run not in chunks:
                chunks.append(run)
    except Exception:
        pass
    # ASCII runs in the binary stream.
    for run in re.findall(rb"[\x20-\x7e\r\n\t]{6,}", raw):
        try:
            text = run.decode("latin-1").strip()
        except Exception:
            continue
        if len(text) >= 6 and text not in chunks:
            chunks.append(text)
    joined = "\n".join(chunks)
    return joined[:limit]


def _looks_text(raw: bytes) -> bool:
    if not raw:
        return True
    sample = raw[:4096]
    if b"\x00" in sample:
        return False
    printable = sum(1 for b in sample if 32 <= b < 127 or b in (9, 10, 13) or b >= 128)
    return printable / max(1, len(sample)) > 0.82


def _strip_xml(xml: str) -> str:
    xml = re.sub(r"<w:tab[^>]*/>", "\t", xml)
    xml = re.sub(r"</w:p>", "\n", xml)
    xml = re.sub(r"<[^>]+>", "", xml)
    xml = (xml.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">")
              .replace("&quot;", '"').replace("&apos;", "'"))
    xml = re.sub(r"[ \t]+\n", "\n", xml)
    xml = re.sub(r"\n{3,}", "\n\n", xml)
    return xml.strip()


def _pdf_extract(path: Path) -> Tuple[str, str, List[Dict[str, str]]]:
    try:
        import pymupdf as fitz
    except Exception:
        try:
            import fitz  # type: ignore
        except Exception:
            return "", "未安装 PyMuPDF，无法读取 PDF；文件已保存，可截图后发送。", []
    images: List[Dict[str, str]] = []
    try:
        doc = fitz.open(str(path))
        pages = []
        empty = []
        for index in range(min(len(doc), 240)):
            text = doc[index].get_text("text") or ""
            if len(text.strip()) < 40:
                empty.append(index)
            pages.append(text)
        text = "\n\n".join(pages).strip()
        note = ""
        if len(text.strip()) < 120 and empty:
            note = "PDF 没有可用文字层（可能是扫描版）。已附上前几页图片，可交给视觉模型阅读。"
        # Render up to MAX_IMAGES pages: prefer empty (scanned) pages first.
        order = (empty + [i for i in range(min(len(doc), 6)) if i not in empty])[:MAX_IMAGES]
        for index in order:
            page = doc[index]
            scale = 1.8
            pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
            data = pix.tobytes("png")
            images.append({"mime": "image/png",
                           "base64": base64.b64encode(data).decode("ascii"),
                           "name": f"{path.name}#page{index + 1}"})
        doc.close()
        if len(text.strip()) < 120 and images:
            ocr_parts = []
            for image in images[:2]:
                try:
                    raw = base64.b64decode(image.get("base64") or "")
                    ocr_text, _ocr_note = _ocr_image(raw)
                except Exception:
                    ocr_text = ""
                if ocr_text:
                    ocr_parts.append(ocr_text)
            if ocr_parts:
                text = "\n\n".join(ocr_parts)
                note = "PDF 无文字层，已用本地 OCR 识别前几页；同时附页面图片供视觉模型查看。"
        return text, note, images
    except Exception as exc:  # noqa: BLE001
        return "", f"PDF 解析失败：{exc}", []


def _docx_extract(path: Path) -> Tuple[str, str]:
    try:
        import docx  # type: ignore
    except Exception:
        docx = None
    if docx is not None:
        try:
            document = docx.Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables[:20]:
                for row in table.rows[:60]:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts), ""
        except Exception:
            pass
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
        return _strip_xml(xml), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"DOCX 解析失败：{exc}"


def _xlsx_extract(path: Path) -> Tuple[str, str]:
    try:
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        chunks = []
        for sheet in wb.worksheets[:12]:
            chunks.append(f"### 工作表：{sheet.title}")
            for row_i, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_i >= 300:
                    break
                cells = [str(c) for c in row[:40] if c is not None]
                if cells:
                    chunks.append(" | ".join(cells))
        wb.close()
        return "\n".join(chunks), ""
    except Exception:
        pass
    try:
        with zipfile.ZipFile(path) as archive:
            shared = []
            if "xl/sharedStrings.xml" in archive.namelist():
                xml = archive.read("xl/sharedStrings.xml").decode("utf-8", errors="ignore")
                shared = re.findall(r"<t[^>]*>(.*?)</t>", xml, re.S)
            rows = []
            for name in archive.namelist():
                if not name.startswith("xl/worksheets/sheet"):
                    continue
                xml = archive.read(name).decode("utf-8", errors="ignore")
                for row in re.findall(r"<row[^>]*>(.*?)</row>", xml, re.S)[:300]:
                    vals = []
                    for cell in re.findall(r"<c[^>]*?(?:t=\"s\")?[^>]*>(.*?)</c>", row, re.S):
                        value = re.search(r"<v>(.*?)</v>", cell)
                        if value:
                            idx = int(value.group(1))
                            vals.append(shared[idx] if 0 <= idx < len(shared) else value.group(1))
                    if vals:
                        rows.append(" | ".join(vals))
            return "\n".join(rows), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"XLSX 解析失败：{exc}"


def _pptx_extract(path: Path) -> Tuple[str, str]:
    try:
        from pptx import Presentation  # type: ignore
        presentation = Presentation(str(path))
        texts = []
        for index, slide in enumerate(presentation.slides[:80], start=1):
            parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows[:30]:
                        cells = [c.text.strip() for c in row.cells if c.text.strip()]
                        if cells:
                            parts.append(" | ".join(cells))
            if parts:
                texts.append(f"## 幻灯片 {index}\n" + "\n".join(parts))
        return "\n\n".join(texts), ""
    except Exception:
        pass
    try:
        with zipfile.ZipFile(path) as archive:
            texts = []
            names = sorted(n for n in archive.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n))
            for name in names[:80]:
                xml = archive.read(name).decode("utf-8", errors="ignore")
                parts = re.findall(r"<a:t>(.*?)</a:t>", xml, re.S)
                if parts:
                    texts.append(" ".join(parts))
        return "\n\n".join(texts), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"PPTX 解析失败：{exc}"


def _ipynb_extract(path: Path) -> Tuple[str, str]:
    try:
        data = json.loads(path.read_text(encoding="utf-8", errors="replace"))
        chunks = []
        for cell in data.get("cells", [])[:300]:
            source = cell.get("source", [])
            if isinstance(source, list):
                source = "".join(source)
            chunks.append(f"[{cell.get('cell_type', 'code')}] {source}")
        return "\n\n".join(chunks), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"Notebook 解析失败：{exc}"


def _ocr_image(raw: bytes) -> Tuple[str, str]:
    try:
        from rapidocr_onnxruntime import RapidOCR  # type: ignore
        import numpy as np
        from PIL import Image
        image = Image.open(io.BytesIO(raw)).convert("RGB")
        if max(image.size) > MAX_IMAGE_SIDE:
            image.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
        array = np.array(image)
        engine = RapidOCR()
        result, _ = engine(array)
        lines = [str(item[1]).strip() for item in (result or [])
                 if len(item) > 1 and str(item[1]).strip()]
        return "\n".join(lines), ""
    except Exception as exc:  # noqa: BLE001
        return "", f"OCR 不可用：{exc}"


def _image_payload(path: Path, raw: bytes) -> Tuple[str, str, List[Dict[str, str]]]:
    text, note = _ocr_image(raw)
    mime = "image/png"
    suffix = path.suffix.lower()
    if suffix in (".jpg", ".jpeg"):
        mime = "image/jpeg"
    elif suffix == ".webp":
        mime = "image/webp"
    try:
        from PIL import Image
        image = Image.open(io.BytesIO(raw))
        if max(image.size) > MAX_IMAGE_SIDE:
            image.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
            buffer = io.BytesIO()
            image.convert("RGB").save(buffer, format="JPEG", quality=88)
            data = buffer.getvalue()
            mime = "image/jpeg"
        else:
            data = raw
    except Exception:
        data = raw
    info = [{"mime": mime, "base64": base64.b64encode(data).decode("ascii"),
             "name": path.name}]
    if not text:
        note = "图片已保存；未识别出文字。若模型支持视觉，可直接查看原图。"
    return text, note, info


def _archive_extract(path: Path) -> Tuple[str, str]:
    try:
        if path.suffix.lower() == ".zip":
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()[:300]
            return "\n".join(names), f"压缩包内共 {len(names)} 个文件（仅列出文件名）。"
        return "", "该压缩格式未解包；文件已保存，可解压后单独上传。"
    except Exception as exc:  # noqa: BLE001
        return "", f"压缩包读取失败：{exc}"


def extract(path: Path, original_name: str = "") -> Dict[str, Any]:
    suffix = path.suffix.lower()
    kind = "text"
    note = ""
    text = ""
    images: List[Dict[str, str]] = []
    raw = b""
    try:
        raw = path.read_bytes()
    except Exception:
        raw = b""

    if suffix in TEXT_SUFFIXES or (not suffix and _looks_text(raw)):
        text = _decode_text(raw)
        kind = "code" if suffix in {
            ".py", ".m", ".c", ".h", ".cpp", ".js", ".ts", ".java", ".v", ".sv",
            ".vhd", ".vhdl", ".rs", ".go", ".sql",
        } else "text"
    elif suffix == ".pdf":
        kind = "pdf"
        text, note, images = _pdf_extract(path)
    elif suffix == ".docx":
        kind = "docx"
        text, note = _docx_extract(path)
    elif suffix in (".xlsx", ".xlsm"):
        kind = "xlsx"
        text, note = _xlsx_extract(path)
    elif suffix == ".pptx":
        kind = "pptx"
        text, note = _pptx_extract(path)
    elif suffix == ".ipynb":
        kind = "notebook"
        text, note = _ipynb_extract(path)
    elif suffix in IMAGE_SUFFIXES:
        kind = "image"
        text, note, images = _image_payload(path, raw)
    elif suffix in AUDIO_SUFFIXES:
        kind = "audio"
        note = "音频已保存。当前接口只发送文字与图片；如需分析音频，请先转写为文本。"
    elif suffix in VIDEO_SUFFIXES:
        kind = "video"
        note = "视频已保存。请截取关键帧为图片后发送，或先转写音轨。"
    elif suffix in ARCHIVE_SUFFIXES:
        kind = "archive"
        text, note = _archive_extract(path)
    elif suffix in LEGACY_OFFICE:
        kind = "legacy_office"
        text = _ole_strings(raw)
        note = ("旧版二进制 Office 文件（.doc/.ppt/.xls）：已尽量提取可读文本；"
                "如排版/公式丢失，可另存为 DOCX/PPTX/XLSX 或 PDF 后重新上传。")
        if not text:
            note = ("旧版二进制 Office 文件（.doc/.ppt/.xls）已保存，但未提取到可读文本；"
                    "建议另存为 DOCX/PPTX/XLSX 或 PDF 后重新上传。")
    elif _looks_text(raw):
        kind = "text"
        text = _decode_text(raw)
    else:
        kind = "binary"
        note = (f"二进制附件已保存（{len(raw) / 1024:.1f} KB）。"
                "模型只能收到文件元数据；如需分析内容，请转换为文本、PDF 或图片。")

    text = re.sub(r"\x00", " ", text or "")
    truncated = len(text) > MAX_CHARS
    if truncated:
        text = text[:MAX_CHARS] + "\n\n…（内容过长已截断）"
    preview = re.sub(r"\s+", " ", text).strip()[:260]
    return {
        "kind": kind,
        "text": text.strip(),
        "chars": len(text),
        "truncated": truncated,
        "note": note,
        "preview": preview,
        "images": images[:MAX_IMAGES],
        "image_count": len(images),
    }


def save_upload(filename: str, content_b64: str) -> Dict[str, Any]:
    config.ensure_dirs()
    original = Path(str(filename or "attachment")).name
    if not original:
        original = "attachment"
    try:
        raw = base64.b64decode(content_b64, validate=True)
    except (binascii.Error, ValueError) as exc:
        raise ValueError("文件内容不是有效的 Base64 数据") from exc
    if len(raw) > MAX_BYTES:
        raise ValueError(f"文件过大（{len(raw) / 1024 / 1024:.1f} MB），上限 60 MB")
    if not raw:
        raise ValueError("文件为空")
    safe = re.sub(r"[^\w\u4e00-\u9fff.\-]+", "_", original)[:100] or "attachment"
    stored = f"{uuid.uuid4().hex[:10]}_{safe}"
    path = config.UPLOAD_DIR / stored
    path.write_bytes(raw)
    info = extract(path, original)
    info.update({
        "ok": True,
        "id": path.name,
        "name": original,
        "size": len(raw),
        "path": str(path),
    })
    _INFO_CACHE[path.name] = info
    _CACHE_ORDER.append(path.name)
    while len(_CACHE_ORDER) > 32:
        old = _CACHE_ORDER.pop(0)
        _INFO_CACHE.pop(old, None)
    return info


def save_text(text: str, name: str = "") -> Dict[str, Any]:
    """Save clipboard text as a text attachment."""
    data = (text or "").encode("utf-8")
    if not data:
        raise ValueError("剪贴板文本为空")
    filename = name or f"clipboard_{uuid.uuid4().hex[:6]}.txt"
    return save_upload(filename, base64.b64encode(data).decode("ascii"))


def load_attachment(attachment_id: str, name: str = "") -> Dict[str, Any]:
    config.ensure_dirs()
    safe = Path(str(attachment_id)).name
    path = config.UPLOAD_DIR / safe
    if not path.exists():
        return {"ok": False, "name": name or safe, "text": "", "images": [],
                "note": "附件不存在或已被清理。"}
    cached = _INFO_CACHE.get(safe)
    if cached:
        if name and not cached.get("name"):
            cached["name"] = name
        return cached
    info = extract(path, name or safe)
    info.update({"ok": True, "id": safe, "name": name or safe,
                 "size": path.stat().st_size})
    _INFO_CACHE[safe] = info
    _CACHE_ORDER.append(safe)
    while len(_CACHE_ORDER) > 32:
        old = _CACHE_ORDER.pop(0)
        _INFO_CACHE.pop(old, None)
    return info


def attachment_context(items: list, max_total: int = 60000) -> Tuple[str, list]:
    """Build a text block for the model and a compact UI description."""
    blocks = []
    used = 0
    ui = []
    for item in items or []:
        if not isinstance(item, dict):
            continue
        info = item
        if not info.get("text") and info.get("id"):
            info = load_attachment(str(info["id"]), str(info.get("name", "")))
        text = info.get("text") or ""
        name = info.get("name") or "附件"
        note = info.get("note") or ""
        if text and used < max_total:
            remaining = max_total - used
            chunk = text[:remaining]
            blocks.append(f"### 附件：{name}（{info.get('kind', 'file')}）\n{chunk}")
            used += len(chunk)
        elif note:
            blocks.append(f"### 附件：{name}\n（{note}）")
        ui.append({
            "id": info.get("id"),
            "name": name,
            "kind": info.get("kind"),
            "chars": info.get("chars", len(text)),
            "preview": info.get("preview", ""),
            "note": note,
            "truncated": info.get("truncated", False),
            "image_count": info.get("image_count", 0),
        })
    return "\n\n".join(blocks), ui


def vision_images(items: list, limit: int = 4) -> List[Dict[str, str]]:
    """Return data URLs for image/PDF attachments (vision models only)."""
    out: List[Dict[str, str]] = []
    for item in items or []:
        if not isinstance(item, dict):
            continue
        info = item
        if not info.get("images") and info.get("id"):
            info = load_attachment(str(info["id"]), str(info.get("name", "")))
        for image in info.get("images") or []:
            mime = image.get("mime") or "image/png"
            b64 = image.get("base64") or ""
            if b64:
                out.append({"url": f"data:{mime};base64,{b64}",
                            "name": image.get("name") or info.get("name") or "image"})
            if len(out) >= limit:
                return out
    return out


def openai_image_blocks(items: list, limit: int = 4) -> List[Dict[str, Any]]:
    return [{"type": "image_url", "image_url": {"url": image["url"]}}
            for image in vision_images(items, limit)]


def anthropic_image_blocks(items: list, limit: int = 4) -> List[Dict[str, Any]]:
    blocks = []
    for image in vision_images(items, limit):
        url = image["url"]
        match = re.match(r"data:([^;]+);base64,(.*)", url, re.S)
        if match:
            blocks.append({
                "type": "image",
                "source": {"type": "base64", "media_type": match.group(1),
                           "data": match.group(2)},
            })
    return blocks
