"""Build an editable DOCX introduction (companion to the HTML/PDF docs)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"
OUT = ROOT / "APP_Release_Release_Version" / "Reference" / "My_APP_Introduction.docx"
OUT_ALT = ROOT / "APP_Release_Release_Version" / "Reference" / "My_APP_Introduction_V4.docx"

BLUE = (0x00, 0x40, 0x98)
DARK = (0x1F, 0x29, 0x37)
MUTED = (0x5B, 0x6B, 0x82)
RED = (0x8B, 0x00, 0x00)


def cn(run, size=10.5, bold=False, color=DARK, name="微软雅黑"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(*color)


def para(doc, text, size=10.5, bold=False, color=DARK, space_after=6,
         align=None, first_indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.45
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    cn(r, size, bold, color)
    return p


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(5)
    border = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18" if level == 1 else "12")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), "004098")
    border.append(left)
    p._p.get_or_add_pPr().append(border)
    r = p.add_run(text)
    cn(r, 15 if level == 1 else 12, True, BLUE)
    return p


def bullet(doc, text, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.4
    if bold_prefix:
        r = p.add_run(bold_prefix)
        cn(r, 10.5, True, DARK)
    r = p.add_run(text)
    cn(r, 10.5, False, DARK)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(h)
        cn(r, 9.5, True, (0x0B, 0x3A, 0x5A))
        shade(c, "EEF6FC")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            r = p.add_run(str(val))
            cn(r, 9.2, False, DARK)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def build() -> Path:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    # cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(50)
    r = p.add_run("智配芯生 · ImpedanceMCP")
    cn(r, 30, True, BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("微波阻抗匹配智能计算与仿真软件")
    cn(r, 16, False, MUTED)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("“智绘芯生，AI 赋能” AI 教与学活动周 · 赛道一（AI 智能体）作品说明")
    cn(r, 11.5, False, RED)

    doc.add_paragraph()
    table(doc, ["项目", "信息"], [
        ["作品名称", "智配芯生 · ImpedanceMCP（微波阻抗匹配智能计算与仿真软件）"],
        ["参赛赛道", "赛道一 · AI 智能体开发"],
        ["团队名称", "智配芯生队"],
        ["团队成员", "卢瑞林（负责人）、卢柯吉、林芷伊、张昀蕾"],
        ["所在学院", "上海交通大学集成电路学院（信息与电子工程学院）"],
        ["提交日期", "2026-10-26"],
        ["版本", "v4.0.0"],
    ], widths=[3.2, 12.8])

    heading(doc, "一、创作背景")
    para(doc, "“微波技术”课程中，阻抗匹配既是重点也是难点：Smith 圆图抽象，学生很难把“串联沿等电阻圆、"
              "并联沿等电导圆”与具体元件值对应起来；L/T/π 网络公式多、符号容易写错，手算后缺少独立校验；"
              "商用射频软件功能强大但学习门槛高，且无法用自然语言提问；大模型能讲理论，却不能可靠地算出"
              "某个具体负载的匹配电容。", first_indent=True)
    para(doc, "团队此前已用 MATLAB App Designer 开发了 V3.0 匹配软件（Impedance_Matching.exe），包含五种拓扑、"
              "Smith 圆图动画和耐压分析。但它仍然是一个“点按钮”的工具，不能与 AI 智能体协作。"
              "因此我们在 V3.0 基础上做了 V4.0：保留原 MATLAB 内核，增加 Python/MATLAB 双计算引擎、"
              "MCP 标准接口、DeepSeek 智能体与现代化可视化界面，让软件既能“自己算”，也能“被 AI 调用”。",
         first_indent=True)

    heading(doc, "二、作品简介")
    para(doc, "智配芯生 · ImpedanceMCP 是一套“AI 智能体 + 微波仿真工具”的桌面软件：")
    for b in [
        ("会算：", "五种拓扑的解析设计 + 完整网络节点法复核；"),
        ("会画：", "Smith 圆图逐步动画、VSWR 带宽、±5% 灵敏度、耐压柱状图；"),
        ("会讲：", "DeepSeek 调用本地工具，用微波工程语言解释结果；"),
        ("会被调用：", "15 个 MCP 工具，ZCode / Claude / Cursor 等客户端可直接使用；"),
        ("会交付：", "一键导出 HTML / Markdown 报告。"),
    ]:
        bullet(doc, b[1], b[0])

    if (ASSETS / "screenshot_design.png").exists():
        doc.add_picture(str(ASSETS / "screenshot_design.png"), width=Cm(16.0))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run("图 1  ImpedanceMCP 匹配设计界面（Smith 圆图、步骤、带宽与灵敏度）")
        cn(r, 9, False, MUTED)

    heading(doc, "三、创新点")
    points = [
        ("1. 把“软件”变成“MCP 服务”",
         "把匹配计算、扫描、复核、耐压、报告封装成 15 个 MCP 工具，任何智能体都能调用。"
         "这不是让大模型背公式，而是让大模型调用真实计算引擎，从根源上减少幻觉。"),
        ("2. 解析解 + 节点法双重校验",
         "每个设计先用拓扑公式求元件值，再把整个网络展开成节点矩阵，独立计算 Zin/VSWR/S11。"
         "300 组随机用例 VSWR 全部 1.000000；MATLAB R2025b 交叉验证 5/5 通过。"
         "同时修复了 V3.0 L 型（右）正向扫描的共轭符号错误。"),
        ("3. 双引擎 + 经典 exe 外壳",
         "Python 内核毫秒级、无 MATLAB 也能用；MATLAB R2025b 一键复核；原版 Impedance_Matching.exe 保留教学动画，一键启动。"),
        ("4. 面向教学的 Smith 圆图动画",
         "三段/四段轨迹逐段绘制，鼠标悬停显示 Z/Γ/VSWR，步骤卡片与轨迹联动，带宽、灵敏度、耐压同屏呈现。"),
        ("5. 工程化的可复现验证",
         "8 项引擎回归测试、11 项 MCP 协议测试、浏览器逐页视觉验收；报告与界面使用同一组精确元件值。"),
    ]
    for title, body in points:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        cn(r, 11, True, BLUE)
        para(doc, body, size=10.3, first_indent=True)

    heading(doc, "四、功能清单")
    table(doc, ["模块", "功能"], [
        ["AI 智能体", "DeepSeek 对话、工具调用、SSE 流式、附件解析、设计卡片、上下文"],
        ["匹配设计", "L 型（左/右）、T 型、π 型Ⅰ/Ⅱ、自动选型、自动 L/C3"],
        ["Smith 圆图", "网格、轨迹动画、播放/暂停/重置/速度、悬停读数"],
        ["工程分析", "VSWR 带宽、±5% 灵敏度、Q 值、风险提示、拓扑对比"],
        ["正向扫描", "可匹配 R/X 范围、点云、扫描结论"],
        ["耐压分析", "峰值/RMS 电压电流、安全系数、柱状图、建议"],
        ["报告", "HTML / Markdown，含圆图、带宽、灵敏度、风险"],
        ["MCP", "15 工具、3 资源、3 提示模板、stdio/HTTP"],
        ["设置", "交大/DeepSeek 预设、模型、Key、引擎、MCP 配置复制"],
        ["经典内核", "启动原版 exe、MATLAB 复核"],
    ], widths=[3.2, 12.8])

    heading(doc, "五、操作指南（5 分钟演示）")
    steps = [
        ("步骤 1：启动", "pip install numpy pypdf openpyxl Pillow，然后双击 start_app.bat；浏览器打开 http://127.0.0.1:8765/。"),
        ("步骤 2：AI 对话", "输入：“在 2.4 GHz 把 20+j5 Ω 匹配到 50 Ω，给出元件值、VSWR 和带宽。”"
                          "智能体调用 design_matching_network，返回 L=1.939 nH、C1=1.624 pF、C2=6.803 pF，VSWR=1.0000，带宽约 57%。"),
        ("步骤 3：打开设计页", "点击设计卡片的“在设计中打开”，查看 Smith 圆图动画、步骤、公式、带宽、灵敏度；点击“MATLAB 复核”做独立复算（约 15–20 秒）。"),
        ("步骤 4：扫描与对比", "点击“扫描范围”，查看 L 型（右）可匹配的 R/X 区域；自动比较五种拓扑，查看带宽/Q；点击“应用”切换方案。"),
        ("步骤 5：耐压", "进入“耐压分析”→“从当前设计填入”→“开始耐压分析”；10 W 下输入峰值电压 31.6 V，C1 峰值 31.6 V，L+C2 19.5 V。"),
        ("步骤 6：报告", "点击“导出报告”→ 浏览器新标签页打开 HTML → Ctrl+P 保存 PDF。"),
        ("步骤 7：MCP 演示", "在 ZCode 中说：“用 impedance-matching 工具设计 100−j30Ω、2.4GHz 的匹配，比较 L 型和 T 型。”"),
    ]
    for title, body in steps:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(title)
        cn(r, 10.8, True, BLUE)
        para(doc, body, size=10.2)

    if (ASSETS / "screenshot_agent.png").exists():
        doc.add_picture(str(ASSETS / "screenshot_agent.png"), width=Cm(16.0))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run("图 2  AI 智能体调用本地工具并生成设计卡片")
        cn(r, 9, False, MUTED)

    heading(doc, "六、技术实现")
    para(doc, "计算：五种拓扑解析解；2/3 节点导纳矩阵，源 50Ω Thevenin，负载 ZL；Zin=Vs/Iin；"
              "带宽 ±60% 扫频 241 点；灵敏度元件 ±5% 重算；耐压由节点电压直接算元件应力。", first_indent=True)
    para(doc, "MCP：JSON-RPC 2.0 over stdio；initialize / tools/list / tools/call / resources / prompts；"
              "15 个工具、3 个资源、3 个提示；本地 HTTP 模式；日志写 stderr，stdout 只输出协议 JSON。", first_indent=True)
    para(doc, "智能体：OpenAI 兼容 DeepSeek 接口；工具循环最多 5 轮；系统提示要求“先算后说”；"
              "附件解析（PDF/DOCX/XLSX/PPTX/Notebook）；SSE 流式事件。", first_indent=True)
    para(doc, "前端：原生 HTML/CSS/JS，无构建步骤；Canvas 2D Smith 圆图；响应式布局；"
              "深色工程控制台主题；报告使用浅色 A4 打印样式。", first_indent=True)

    heading(doc, "七、测试与验证")
    table(doc, ["测试", "结果"], [
        ["Python 引擎回归", "8/8 通过"],
        ["300 组随机设计", "VSWR 全部 1.000000"],
        ["MATLAB R2025b 交叉", "5/5 通过，Zin 偏差 0.0"],
        ["MCP 协议", "11/11 通过"],
        ["浏览器视觉", "AI/设计/分析/耐压/设置/关于/报告逐页检查"],
        ["报告一致性", "报告与界面使用同一组精确元件值"],
    ], widths=[7.0, 9.0])
    para(doc, "复现命令：python -m unittest app.tests.test_engine -v；"
              "python -m unittest app.tests.test_mcp_protocol -v。", size=9.5, color=MUTED)

    heading(doc, "八、团队分工")
    table(doc, ["成员", "分工"], [
        ["卢瑞林", "项目负责人；总体架构、Python 计算内核、MCP 协议、DeepSeek 智能体、集成与发布"],
        ["卢柯吉", "MATLAB 计算内核、mw_design.m、MATLAB R2025b 交叉验证、原版 exe 维护"],
        ["林芷伊", "交互设计、Smith 圆图与图表可视化、UI 样式、用户手册与演示材料"],
        ["张昀蕾", "测试用例、MCP/浏览器验收、演示脚本、用户研究与反馈整理"],
    ], widths=[3.2, 12.8])

    heading(doc, "九、总结")
    para(doc, "智配芯生 · ImpedanceMCP 不是“给旧软件套一个聊天框”，而是把微波匹配计算重新组织成可被智能体调用的能力："
              "对学习者，圆图动画 + 对话解释降低理解门槛；对工程师，解析解 + 节点法 + 带宽/容差/耐压可直接用于方案评估；"
              "对 AI 生态，15 个 MCP 工具让大模型“算得准、说得清”；对教学，保留原 MATLAB 软件，形成“经典内核 + 现代外壳”的渐进式升级。",
         first_indent=True)
    para(doc, "我们希望说明：AI 不是替代专业工具，而是让专业工具更容易被使用、被验证、被复用。", first_indent=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("仓库：https://github.com/RuilinLu/Impedance_Matching_Calculation\n"
                  "版本 v4.0.0 ｜ 2026-10-26 ｜ 团队：卢瑞林、卢柯吉、林芷伊、张昀蕾\n"
                  "上海交通大学集成电路学院（信息与电子工程学院）")
    cn(r, 9, False, MUTED)

    for target in (OUT, OUT_ALT):
        try:
            doc.save(str(target))
            print("saved", target)
            return target
        except PermissionError as exc:
            print("locked:", target, exc)
    return OUT_ALT


if __name__ == "__main__":
    build()
